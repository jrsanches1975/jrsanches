#!/usr/bin/env python3
"""
Agente de determinação de produtos para monitorar no Google Shopping.

Decide QUAIS produtos da Joie merecem monitoramento no Google Shopping, usando
evidência REAL de dois canais — nunca por achismo:

  1. Anúncios do site (Google Ads): campanhas com o padrão de nome
     "Shopping - <produto>" são confirmação DIRETA de que o produto já roda em
     Shopping Ads. Ex. real visto na conta da Joie: campanha "Shopping - Colageno",
     R$ 422,78 de gasto, 1.230 cliques — ou seja, Colágeno já está em Shopping
     mesmo sem estar cadastrado em produtos_monitorados/produtos_candidatos_manual.
  2. Anúncios nos marketplaces (Mercado Livre): o NOSSO próprio anúncio
     (snapshot_proprio, o mesmo que alimenta o Radar de Posição) aparecendo de
     fato na busca é sinal de que o produto está no catálogo ativo — mesmo sem
     campanha Shopping dedicada ainda, é candidato a ganhar uma.

Cruza os dois com o que já está em produtos_monitorados/produtos_candidatos_manual
e classifica cada produto:
  - "já monitorado no Shopping" — tem campanha Shopping real detectada
  - "candidato — vende no ML mas sem Shopping" — tem anúncio próprio no ML,
    sem campanha Shopping (oportunidade)
  - "ACHADO fora do config" — tem campanha Shopping real, mas nem está
    cadastrado como produto no config (o caso real do Colágeno)
  - "sem evidência" — nem campanha Shopping nem anúncio próprio confirmado
    (não promove sozinho, só sinaliza pra revisão manual)

NÃO chama nenhuma API sozinho — os JSONs de entrada são coletas já feitas
(mesmo formato de own_performance.py --out e do snapshot_proprio do Radar em
war_room.py). A coleta de CONCORRENTES no Google Shopping (quem mais aparece
lá pra esses produtos) ainda depende de um ator do Apify pra Google Shopping,
que este script NÃO inclui — nenhum foi testado ao vivo ainda (mesmo status
BETA de meta_ads.py/google_ads_transparency.py). Esse é o próximo passo
natural, uma vez que este agente já diz QUAIS produtos monitorar.

Uso:
    python descoberta_produtos_shopping.py --config config.json \
        --own-performance-json ../outputs/own-performance-por-produto.json \
        --out ../outputs/descoberta-shopping.xlsx \
        --export-json ../outputs/descoberta-shopping.json

    # se já tiver o snapshot do NOSSO anúncio no Mercado Livre (vem do Radar):
    python descoberta_produtos_shopping.py --config config.json \
        --own-performance-json ../outputs/own-performance-por-produto.json \
        --snapshot-proprio-json ../outputs/radar-proprio-ml.json \
        --out ../outputs/descoberta-shopping.xlsx
"""
import argparse
import re
import sys

from apify_common import load_json, norm, save_json

PADRAO_SHOPPING = re.compile(r"^shopping\s*-\s*(.+?)\.?$", re.IGNORECASE)


def detectar_campanhas_shopping(por_campanha):
    """{nome_bruto_extraido: {campanha, spend, clicks}} pra cada campanha real do
    Google Ads cujo nome bate o padrão 'Shopping - X'. Só considera o conector
    google_ads (a mesma campanha aparece também via googleanalytics4, sem
    spend/clicks — não é a fonte de verdade do gasto)."""
    achados = {}
    for row in por_campanha or []:
        if row.get("conector") != "google_ads":
            continue
        campanha = (row.get("campanha") or "").strip()
        m = PADRAO_SHOPPING.match(campanha)
        if not m:
            continue
        nome_bruto = m.group(1).strip()
        achados[nome_bruto] = {
            "campanha": campanha,
            "spend": row.get("spend"),
            "clicks": row.get("clicks"),
        }
    return achados


def casar_produto_config(nome_bruto, produtos_cfg):
    """Tenta casar o nome extraído de 'Shopping - X' contra os produtos já
    configurados — por nome ou pelas substrings de campanha (mesma lógica de
    match_produto em own_performance.py: substring, case-insensitive, sensível a
    acento). Retorna o nome oficial do produto no config, ou None se não bater
    com nenhum (produto novo, ainda não configurado)."""
    alvo = norm(nome_bruto)
    if not alvo:
        return None
    for p in produtos_cfg:
        termos = [p.get("nome", "")] + p.get("campanhas_google_ads", []) + p.get("campanhas_meta_ads", [])
        for termo in termos:
            t = norm(termo)
            if t and (t in alvo or alvo in t):
                return p["nome"]
    return None


def montar_recomendacoes(campanhas_shopping, produtos_cfg, snapshot_proprio):
    """Uma linha por produto: os já configurados (com ou sem Shopping real) +
    os achados só via campanha Shopping mas ainda sem entrada no config."""
    snapshot_proprio = snapshot_proprio or {}
    recomendacoes = []
    casados = set()

    for p in produtos_cfg:
        nome = p["nome"]
        shopping_real = None
        for nome_bruto, dados in campanhas_shopping.items():
            if casar_produto_config(nome_bruto, [p]) == nome:
                shopping_real = dados
                casados.add(nome_bruto)
                break
        tem_anuncio_ml = nome in snapshot_proprio

        if shopping_real:
            status = "já monitorado no Shopping (campanha real ativa)"
        elif tem_anuncio_ml:
            status = "candidato — vende no ML mas sem campanha Shopping detectada"
        else:
            status = "sem evidência de anúncio confirmada — revisar à mão"

        recomendacoes.append({
            "produto": nome,
            "termo_busca_ml": p.get("termo_busca_ml"),
            "campanha_shopping_real": shopping_real,
            "anuncio_proprio_ml_confirmado": tem_anuncio_ml,
            "status": status,
            "origem": "config",
        })

    for nome_bruto, dados in campanhas_shopping.items():
        if nome_bruto in casados:
            continue
        if casar_produto_config(nome_bruto, produtos_cfg):
            continue  # casou por outro caminho (não deveria, defensivo)
        recomendacoes.append({
            "produto": nome_bruto,
            "termo_busca_ml": None,
            "campanha_shopping_real": dados,
            "anuncio_proprio_ml_confirmado": False,
            "status": "ACHADO: tem campanha Shopping real, mas não está cadastrado "
                      "como produto em produtos_monitorados/produtos_candidatos_manual "
                      "— considerar adicionar",
            "origem": "campanha_shopping",
        })
    return recomendacoes


def write_xlsx(recomendacoes, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    NAVY = "1F2A44"
    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor=NAVY)

    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos p_ Shopping"

    cols = ["produto", "status", "origem", "termo_busca_ml", "campanha_shopping_real",
            "gasto_shopping", "cliques_shopping", "anuncio_proprio_ml_confirmado"]
    ws.append([c.upper() for c in cols])
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hfont
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    recomendacoes = sorted(
        recomendacoes,
        key=lambda r: (0 if r["campanha_shopping_real"] else 1,
                       0 if r["anuncio_proprio_ml_confirmado"] else 1,
                       r["produto"]),
    )
    for r in recomendacoes:
        camp = r["campanha_shopping_real"] or {}
        ws.append([
            r["produto"],
            r["status"],
            r["origem"],
            r.get("termo_busca_ml") or "—",
            camp.get("campanha") or "—",
            round(camp["spend"], 2) if camp.get("spend") is not None else "—",
            camp.get("clicks") if camp.get("clicks") is not None else "—",
            "sim" if r["anuncio_proprio_ml_confirmado"] else "não/não verificado",
        ])
    larguras = [22, 46, 16, 26, 26, 14, 14, 26]
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    import os
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(
        description="Determina quais produtos da Joie monitorar no Google Shopping, "
                    "com base em campanhas Shopping reais e anúncio próprio no ML.")
    ap.add_argument("--config", required=True)
    ap.add_argument("--own-performance-json", required=True,
                     help="saída de own_performance.py (tem por_campanha, com as campanhas "
                          "reais do Google Ads — é aí que 'Shopping - X' é detectado)")
    ap.add_argument("--snapshot-proprio-json", default=None,
                     help="opcional: {produto: {...}} do NOSSO anúncio no Mercado Livre "
                          "(mesmo formato do snapshot_proprio interno do Radar). Sem isso, "
                          "'anuncio_proprio_ml_confirmado' fica sempre 'não verificado'.")
    ap.add_argument("--out", default="descoberta-shopping.xlsx")
    ap.add_argument("--export-json", default=None,
                     help="grava a lista de recomendações em JSON, pra consulta ou "
                          "integração futura com war_room.py")
    args = ap.parse_args()

    config = load_json(args.config, {})
    produtos_cfg = config.get("produtos_monitorados", []) + config.get("produtos_candidatos_manual", [])

    own_perf = load_json(args.own_performance_json, {})
    por_campanha = own_perf.get("por_campanha", [])
    if not por_campanha:
        print("AVISO: own-performance-json não tem 'por_campanha' ou veio vazio — "
              "nenhuma campanha Shopping pode ser detectada nesta rodada.", file=sys.stderr)

    snapshot_proprio = load_json(args.snapshot_proprio_json, None) if args.snapshot_proprio_json else None
    if args.snapshot_proprio_json and snapshot_proprio is None:
        print(f"AVISO: não consegui ler {args.snapshot_proprio_json}", file=sys.stderr)

    campanhas_shopping = detectar_campanhas_shopping(por_campanha)
    recomendacoes = montar_recomendacoes(campanhas_shopping, produtos_cfg, snapshot_proprio)

    write_xlsx(recomendacoes, args.out)
    if args.export_json:
        save_json(args.export_json, {"recomendacoes": recomendacoes})
        print(f"OK -> {args.export_json}", file=sys.stderr)

    n_shopping = sum(1 for r in recomendacoes if r["campanha_shopping_real"])
    n_achados = sum(1 for r in recomendacoes if r["origem"] == "campanha_shopping")
    print(f"OK -> {args.out} ({len(recomendacoes)} produto(s) avaliado(s), "
          f"{n_shopping} com campanha Shopping real, {n_achados} achado(s) fora do config)")
    if not snapshot_proprio:
        print("  AVISO: rodado sem --snapshot-proprio-json — 'vende no ML' não foi "
              "verificado nesta rodada, só a parte de campanha Shopping.", file=sys.stderr)


if __name__ == "__main__":
    main()
