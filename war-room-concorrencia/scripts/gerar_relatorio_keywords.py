#!/usr/bin/env python3
"""
Gera o RELATÓRIO COMPLETO de palavras-chave + leilão — a relação de TODAS as
keywords usadas nas campanhas (não só as que caíram de performance, diferente
de keyword_auction.py, que só aponta quedas) com as métricas de leilão de cada
uma (impression share, rank lost, Quality Score, CPC) e os domínios
concorrentes disputando cada campanha (Auction Insight), ordenados por
frequência no período.

Reaproveita 100% do motor de agregação de keyword_auction.py
(agregar_keywords, agregar_dominios_por_campanha) — não reimplementa nada, só
troca "detectar e alertar quedas" por "listar tudo".

Uso:
    python gerar_relatorio_keywords.py --config config.example.json \
        --keywords-json examples/keywords-simulado.json \
        --auction-json examples/auction-domains-simulado.json \
        --out ../outputs/relatorio-keywords.xlsx
"""
import argparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import keyword_auction as ka
from apify_common import load_json

NAVY = "1F2A44"


def style_header(ws, ncols):
    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor=NAVY)
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hfont
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def wrap_body(ws, min_row=2):
    for row in ws.iter_rows(min_row=min_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def fmt(v, casas=2):
    return round(v, casas) if isinstance(v, (int, float)) else "N/D"


def write_xlsx(keywords_agg, dominios_por_campanha, simulado, path):
    wb = Workbook()

    # -- aba 1: todas as keywords --
    ws = wb.active
    ws.title = "Keywords"
    cols = ["campanha", "keyword", "impressions", "clicks", "ctr_%", "cpc_medio",
            "quality_score", "impression_share", "rank_lost_impression_share",
            "first_page_cpc", "top_of_page_cpc", "dominios_no_leilao"]
    ws.append([c.upper() for c in cols])
    style_header(ws, len(cols))

    linhas = sorted(keywords_agg.values(), key=lambda a: -(a["impressions"] or 0))
    for a in linhas:
        ctr_pct = (a["clicks"] / a["impressions"] * 100) if a["impressions"] else None
        doms = dominios_por_campanha.get(a["campanha"], [])
        doms_texto = ", ".join(f"{d} ({n}x)" for d, n in doms[:5]) or "(nenhum capturado)"
        ws.append([
            a["campanha"], a["keyword"], a["impressions"], a["clicks"],
            fmt(ctr_pct, 1),
            fmt(a["cpc_medio"]),
            fmt(a["quality_score_medio"], 1),
            fmt(a["impression_share_medio"] * 100, 1) if a["impression_share_medio"] is not None else "N/D",
            fmt(a["rank_lost_medio"] * 100, 1) if a["rank_lost_medio"] is not None else "N/D",
            fmt(a["first_page_cpc"]) if a["first_page_cpc"] is not None else "N/D (não estimado pelo Google)",
            fmt(a["top_of_page_cpc"]) if a["top_of_page_cpc"] is not None else "N/D (não estimado pelo Google)",
            doms_texto,
        ])
    for i, w in enumerate([24, 26, 12, 10, 9, 11, 13, 15, 20, 16, 16, 46], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wrap_body(ws)
    ws.freeze_panes = "A2"

    # -- aba 2: leilão por campanha (todos os domínios, não só os top 5 da aba 1) --
    ws2 = wb.create_sheet("Leilão por Campanha")
    cols2 = ["campanha", "dominio_concorrente", "aparicoes_no_periodo"]
    ws2.append([c.upper() for c in cols2])
    style_header(ws2, len(cols2))
    for campanha in sorted(dominios_por_campanha):
        for dom, n in dominios_por_campanha[campanha]:
            ws2.append([campanha, dom, n])
    for i, w in enumerate([28, 30, 20], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    wrap_body(ws2)
    ws2.freeze_panes = "A2"

    # -- aba 3: metodologia / limitações --
    ws3 = wb.create_sheet("Metodologia")
    notas = [
        ("Fonte", ""),
        ("Google Ads via Windsor.ai", "impressions/clicks/ctr/cpc/quality_score/search_impression_share/"
         "search_rank_lost_impression_share vêm de uma chamada get_data no connector google_ads. "
         "auction_insight_domain vem de uma chamada SEPARADA — o Google Ads recusa combinar "
         "esse campo com métricas de performance na mesma consulta (\"unsupported metrics\")."),
        ("first_page_cpc / top_of_page_cpc", "Estimativas de lance do próprio Google para aparecer na "
         "1a página/topo. Em termos de baixo volume o Google costuma não popular essas estimativas — "
         "quando vier N/D, é porque o Google não devolveu valor, nunca um número inventado aqui."),
        ("Dados deste relatório", "SIMULADO — ver aviso abaixo." if simulado else
         "REAL — puxado ao vivo da conta Google Ads via Windsor.ai."),
        ("", ""),
        ("Como gerar com dado real", "1) Puxe via MCP do Windsor.ai: get_data(connector=\"google_ads\", "
         "fields=[\"date\",\"campaign\",\"keyword_text\",\"impressions\",\"clicks\",\"ctr\",\"cpc\","
         "\"quality_score\",\"search_impression_share\",\"search_rank_lost_impression_share\","
         "\"first_page_cpc\",\"position_estimates_top_of_page_cpc_micros\"]) → grave em keywords.json. "
         "2) get_data(connector=\"google_ads\", fields=[\"date\",\"campaign\",\"auction_insight_domain\"]) → "
         "grave em auction.json. 3) Rode gerar_relatorio_keywords.py apontando pros dois arquivos."),
        ("Pré-requisito", "A conta Google Ads precisa estar conectada no Windsor.ai (get_connectors deve "
         "mostrar accounts para o connector google_ads) — ver references/fontes-e-limitacoes.md para o "
         "status mais recente dessa conexão."),
    ]
    for a, b in notas:
        ws3.append([a, b])
        if b == "" and a:
            ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, color=NAVY, size=12)
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 100
    wrap_body(ws3, min_row=1)

    if simulado:
        ws4 = wb.create_sheet("⚠ AVISO", 0)
        ws4.sheet_view.showGridLines = False
        ws4.column_dimensions["A"].width = 100
        aviso = [
            "⚠ ESTE RELATÓRIO USA DADO SIMULADO, NÃO REAL.",
            "",
            "A conta Google Ads da Joie ainda não está conectada no Windsor.ai desta integração "
            "(get_connectors só mostra a GA4 — ver references/fontes-e-limitacoes.md, achado de "
            "2026-08-01). Os números de keyword/leilão abaixo (impression share, CPC, Quality Score, "
            "domínios concorrentes) servem só para mostrar o FORMATO do relatório — não são dado real "
            "da conta.",
            "",
            "Assim que o Google Ads estiver conectado (get_connectors mostrando accounts para "
            "google_ads), rode gerar_relatorio_keywords.py com os JSONs reais puxados via Windsor.ai "
            "(ver aba Metodologia) para substituir por dado medido de verdade.",
        ]
        for linha in aviso:
            ws4.append([linha])
        ws4.cell(row=1, column=1).font = Font(bold=True, color="B00020", size=13)
        wrap_body(ws4, min_row=1)

    import os
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description="Relatório completo de keywords + leilão (todas, não só quedas).")
    ap.add_argument("--config", required=True)
    ap.add_argument("--keywords-json", required=True)
    ap.add_argument("--auction-json", required=True)
    ap.add_argument("--out", default="relatorio-keywords.xlsx")
    ap.add_argument("--simulado", action="store_true",
                     help="marca o relatório como dado simulado (adiciona aba de aviso) — use sempre que "
                          "os JSONs de entrada não vieram de uma coleta real via Windsor.ai")
    ap.add_argument("--export-json", default=None,
                     help="grava {keywords, dominios_por_campanha, simulado} em JSON para o war_room.py "
                          "ingerir como aba 'Keywords & Leilão' — ver war_room.py --keywords-relatorio-json")
    args = ap.parse_args()

    config = load_json(args.config, {})
    dominios_proprios = config.get("official_domains", []) or [
        c.get("dominio_site", "") for c in config.get("concorrentes", [])
    ]

    keywords_regs = ka.carregar_registros(args.keywords_json)
    auction_regs = ka.carregar_registros(args.auction_json)

    keywords_agg = ka.agregar_keywords(keywords_regs)
    dominios_por_campanha = ka.agregar_dominios_por_campanha(auction_regs, dominios_proprios)

    write_xlsx(keywords_agg, dominios_por_campanha, args.simulado, args.out)

    if args.export_json:
        from apify_common import save_json
        save_json(args.export_json, {
            "keywords": keywords_agg,
            "dominios_por_campanha": {k: [list(t) for t in v] for k, v in dominios_por_campanha.items()},
            "simulado": args.simulado,
        })
        print(f"OK -> {args.export_json} (JSON p/ war_room.py)")

    print(f"OK -> {args.out} ({len(keywords_agg)} keyword(s), "
          f"{sum(len(v) for v in dominios_por_campanha.values())} registro(s) de leilão)"
          + (" [SIMULADO]" if args.simulado else " [DADO REAL]"))


if __name__ == "__main__":
    main()
