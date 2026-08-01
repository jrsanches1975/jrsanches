#!/usr/bin/env python3
"""
Motor de descoberta e composição de concorrentes — três fontes viram uma lista única
de candidatos, ranqueada por relevância, para você decidir quem promover a
`concorrentes[]` no config principal:

  1. Introdução MANUAL (config["candidatos_concorrentes_manual"]).
  2. Descoberta POR PRODUTO: para cada produto em produtos_monitorados +
     produtos_candidatos_manual, busca o termo no Mercado Livre e — diferente de
     collect_snapshot() (que só vê quem já está cadastrado) — captura TODOS os
     sellers distintos que aparecem, exceto o próprio (official_sellers). Isso é
     "procurar variantes similares e trazer os concorrentes só daquele produto".
  3. AUCTION INSIGHT (Google Ads, via Windsor.ai/keyword_auction.py) — os domínios
     que disputam o leilão da própria conta são concorrentes confirmados por dado
     real, não por busca.

Cada candidato é pontuado por relevância combinando (pesos em config
["pesos_relevancia"], documentados em references/fontes-e-limitacoes.md):
  - preco: proximidade do preço ao seu preco_proprio no produto.
  - autoridade_marca: PROXY via rating médio do seller no Mercado Livre.
  - presenca_ads: nº de anúncios ativos encontrados (se você rodou meta_ads.py/
    google_ads_transparency.py e passar --ativos-ads-json).
  - vendas_marketplace: PROXY via nº de reviews e de listagens do seller.
  - kpis_redes_sociais: SEM FONTE DE DADO REAL integrada ainda — sempre "N/D",
    nunca inventado. O peso configurado fica documentado, mas não entra na conta
    até existir uma fonte real (ver references/fontes-e-limitacoes.md).

Um componente sem dado é removido do somatório e os pesos dos demais são
renormalizados — nunca vira 0 por "falta de dado", vira "não considerado".

Uso:
    python descoberta_concorrentes.py --config config.json --out ../outputs/descoberta.xlsx
    python descoberta_concorrentes.py --config config.example.json \
        --simulate-raw examples/descoberta-simulada.json --out ../outputs/descoberta.xlsx
"""
import argparse
import math
import sys
from collections import defaultdict

from apify_common import apify_run, get_token, load_json, norm
from war_room import ML_ACTOR, match_oficial

MAX_REVIEWS_REFERENCIA = 2000  # teto pra normalizar autoridade/vendas (log-scale)


# --------------------------------------------------------------------- descoberta
def buscar_raw(termo, token, max_items):
    items, err = apify_run(ML_ACTOR, {"searchQuery": termo, "maxItems": max_items}, token)
    if err:
        print(f"  ERRO ao buscar '{termo}': {err}", file=sys.stderr)
    return items or []


def descobrir_por_produto(raw_items, official_sellers):
    """Agrega TODOS os sellers distintos do resultado bruto (exceto o próprio) —
    cada um é um candidato a concorrente ESPECÍFICO daquele produto/variante."""
    por_seller = {}
    for idx, item in enumerate(raw_items):
        seller = item.get("seller") or {}
        nick = seller.get("nickname") if isinstance(seller, dict) else (seller or "")
        if not nick or (official_sellers and match_oficial(nick, official_sellers)):
            continue
        position = idx + 1
        if nick not in por_seller:
            por_seller[nick] = {
                "seller": nick, "title": item.get("title"), "price": item.get("price"),
                "position": position, "reviews": item.get("reviews_count"),
                "rating": item.get("average_rating"), "total_listagens": 1,
                "url": item.get("url"), "origem": ["descoberta_ml"],
            }
        else:
            por_seller[nick]["total_listagens"] += 1
            if position < por_seller[nick]["position"]:
                por_seller[nick]["position"] = position
    return por_seller


def carregar_dominios_auction(path, official_domains):
    """Mesmo formato de registros do keyword_auction.py (date, campaign,
    auction_insight_domain) — aqui agregado globalmente (não por campanha), porque
    a intenção é 'quem disputa a conta inteira', não um produto específico."""
    registros = load_json(path, {}).get("registros", [])
    contagem = defaultdict(int)
    proprios = {norm(d) for d in official_domains}
    for r in registros:
        dom = (r.get("auction_insight_domain") or "").strip()
        if not dom or norm(dom) in proprios:
            continue
        contagem[dom] += 1
    return dict(contagem)


# ---------------------------------------------------------------------- composição
def _termos_manual(c):
    """Todos os textos que identificam um candidato manual, pra casar por substring
    contra o que a descoberta/leilão trazem (mesma marca pode aparecer como seller
    ML, domínio ou nome digitado à mão — sem isso, viraria 3 candidatos separados)."""
    termos = [c.get("nome"), c.get("dominio_site")] + list(c.get("sellers_ml", []))
    return [norm(t) for t in termos if t]


def _acha_manual_por_substring(texto, manuais_termos):
    t = norm(texto)
    for i, termos in enumerate(manuais_termos):
        if any(termo in t or t in termo for termo in termos):
            return i
    return None


def compor_candidatos_produto(produto_nome, descoberta_ml, candidatos_manual):
    candidatos = [{**c, "origem": ["manual"], "produto": produto_nome} for c in candidatos_manual]
    manuais_termos = [_termos_manual(c) for c in candidatos_manual]
    for nick, dados in descoberta_ml.items():
        i = _acha_manual_por_substring(nick, manuais_termos)
        if i is not None:
            candidatos[i]["origem"].append("descoberta_ml")
            candidatos[i].update({k: v for k, v in dados.items() if k not in ("origem",)})
        else:
            candidatos.append({**dados, "nome": nick, "produto": produto_nome})
    return candidatos


def compor_candidatos_globais(dominios_auction, candidatos_manual):
    candidatos = [{**c, "origem": ["manual"]} for c in candidatos_manual]
    manuais_termos = [_termos_manual(c) for c in candidatos_manual]
    for dominio, aparicoes in dominios_auction.items():
        i = _acha_manual_por_substring(dominio, manuais_termos)
        if i is not None:
            candidatos[i]["origem"].append("auction_insight")
            candidatos[i]["aparicoes_leilao"] = aparicoes
        else:
            candidatos.append({"nome": dominio, "dominio_site": dominio,
                                "aparicoes_leilao": aparicoes, "origem": ["auction_insight"]})
    return candidatos


# ------------------------------------------------------------------------ scoring
def _normalizar_log(valor, teto):
    if valor is None or valor <= 0:
        return None
    return min(math.log10(valor + 1) / math.log10(teto + 1), 1.0)


def pontuar(candidato, preco_proprio, ativos_ads_por_nome, pesos):
    componentes = {}

    # preço: proximidade ao preço próprio (quanto mais perto, mais é substituto direto)
    if candidato.get("price") is not None and preco_proprio:
        delta_pct = abs(candidato["price"] - preco_proprio) / preco_proprio * 100
        componentes["preco"] = 1 / (1 + delta_pct / 50)
    else:
        componentes["preco"] = None

    # autoridade de marca: PROXY via rating médio (não é medida real de brand equity)
    if candidato.get("rating") is not None:
        componentes["autoridade_marca"] = max(0.0, min(candidato["rating"] / 5.0, 1.0))
    else:
        componentes["autoridade_marca"] = None

    # presença em ads: nº de anúncios ativos encontrados (meta_ads.py/google_ads_transparency.py)
    nome_norm = norm(candidato.get("nome") or candidato.get("seller") or "")
    n_ads = ativos_ads_por_nome.get(nome_norm)
    componentes["presenca_ads"] = min(n_ads / 10, 1.0) if n_ads is not None else None

    # vendas em marketplace: PROXY via reviews (volume) + total_listagens (catálogo)
    rev_score = _normalizar_log(candidato.get("reviews"), MAX_REVIEWS_REFERENCIA)
    list_score = _normalizar_log(candidato.get("total_listagens"), 50)
    partes = [s for s in (rev_score, list_score) if s is not None]
    componentes["vendas_marketplace"] = sum(partes) / len(partes) if partes else None

    # KPIs de redes sociais: sem fonte real integrada — sempre N/D, nunca inventado
    componentes["kpis_redes_sociais"] = None

    disponiveis = {k: v for k, v in componentes.items() if v is not None}
    peso_disponivel = sum(pesos.get(k, 0) for k in disponiveis)
    if peso_disponivel > 0:
        score = sum(pesos.get(k, 0) * v for k, v in disponiveis.items()) / peso_disponivel
    else:
        score = None

    return {
        "score_final": score,
        "componentes": componentes,
        "componentes_indisponiveis": [k for k, v in componentes.items() if v is None],
    }


# -------------------------------------------------------------------------- xlsx
def write_xlsx(por_produto_pontuado, globais_pontuado, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    NAVY = "1F2A44"
    hfont = Font(bold=True, color="FFFFFF", size=11)
    hfill = PatternFill("solid", fgColor=NAVY)

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = hfont
            cell.fill = hfill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cols = ["produto", "candidato", "score", "preco", "autoridade_marca", "presenca_ads",
            "vendas_marketplace", "kpis_redes_sociais", "origem", "titulo/variante", "componentes_sem_dado"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Candidatos por Produto"
    ws.append([c.upper() for c in cols])
    style_header(ws, len(cols))
    linhas = []
    for produto, candidatos in por_produto_pontuado.items():
        for cand, pont in candidatos:
            linhas.append((produto, cand, pont))
    linhas.sort(key=lambda t: -(t[2]["score_final"] or -1))
    for produto, cand, pont in linhas:
        comp = pont["componentes"]
        ws.append([
            produto, cand.get("nome") or cand.get("seller"),
            round(pont["score_final"], 3) if pont["score_final"] is not None else "N/D",
            round(comp["preco"], 2) if comp["preco"] is not None else "N/D",
            round(comp["autoridade_marca"], 2) if comp["autoridade_marca"] is not None else "N/D",
            round(comp["presenca_ads"], 2) if comp["presenca_ads"] is not None else "N/D",
            round(comp["vendas_marketplace"], 2) if comp["vendas_marketplace"] is not None else "N/D",
            "N/D (sem fonte real integrada)",
            ", ".join(cand.get("origem", [])),
            cand.get("title") or cand.get("dominio_site") or "",
            ", ".join(pont["componentes_indisponiveis"]),
        ])
    for i, w in enumerate([20, 24, 9, 9, 15, 13, 16, 26, 22, 40, 34], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Candidatos Globais (Leilão)")
    cols2 = ["candidato", "score", "preco", "autoridade_marca", "presenca_ads",
             "vendas_marketplace", "aparicoes_leilao", "origem", "componentes_sem_dado"]
    ws2.append([c.upper() for c in cols2])
    style_header(ws2, len(cols2))
    globais_ordenados = sorted(globais_pontuado, key=lambda t: -(t[1]["score_final"] or -1))
    for cand, pont in globais_ordenados:
        comp = pont["componentes"]
        ws2.append([
            cand.get("nome"),
            round(pont["score_final"], 3) if pont["score_final"] is not None else "N/D",
            round(comp["preco"], 2) if comp["preco"] is not None else "N/D",
            round(comp["autoridade_marca"], 2) if comp["autoridade_marca"] is not None else "N/D",
            round(comp["presenca_ads"], 2) if comp["presenca_ads"] is not None else "N/D",
            round(comp["vendas_marketplace"], 2) if comp["vendas_marketplace"] is not None else "N/D",
            cand.get("aparicoes_leilao", "—"),
            ", ".join(cand.get("origem", [])),
            ", ".join(pont["componentes_indisponiveis"]),
        ])
    for i, w in enumerate([24, 9, 9, 15, 13, 16, 16, 22, 34], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Metodologia")
    notas = [
        ("Como o score é calculado", ""),
        ("Preço", "Proximidade ao preco_proprio do produto — 1/(1+delta_%/50). Quanto mais perto do seu preço, mais substituto direto."),
        ("Autoridade de marca", "PROXY via rating médio do seller no Mercado Livre (0-5 normalizado). NÃO é medida real de brand equity."),
        ("Presença em ads", "Nº de anúncios ativos encontrados via meta_ads.py/google_ads_transparency.py (--ativos-ads-json). Sem esse arquivo, fica N/D."),
        ("Vendas em marketplace", "PROXY via nº de reviews (escala log) e nº de listagens do seller — não é dado de venda real."),
        ("KPIs de redes sociais", "SEM FONTE DE DADO REAL integrada nesta versão. Sempre N/D — nunca inventado. Precisaria de uma integração nova (ex. Instagram/TikTok API)."),
        ("", ""),
        ("Renormalização", "Um componente sem dado sai do somatório e os pesos dos demais são redistribuídos proporcionalmente — nunca vira 0 por falta de dado."),
    ]
    for a, b in notas:
        ws3.append([a, b])
        if b == "" and a:
            ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True, color=NAVY, size=12)
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["B"].width = 100
    for row in ws3.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(path)


# ------------------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description="Motor de descoberta e composição de concorrentes.")
    ap.add_argument("--config", required=True)
    ap.add_argument("--out", default="descoberta-concorrentes.xlsx")
    ap.add_argument("--auction-json", default=None, help="registros de auction_insight_domain (mesmo formato do keyword_auction.py)")
    ap.add_argument("--ativos-ads-json", default=None, help="JSON {nome_candidato: nº_anuncios_ativos}")
    ap.add_argument("--per-produto", type=int, default=20)
    ap.add_argument("--simulate-raw", default=None,
                     help="JSON {produto: [item_bruto, ...]} pra testar sem coleta real (ver scripts/examples/)")
    ap.add_argument("--token", default=None)
    args = ap.parse_args()

    config = load_json(args.config, {})
    official_sellers = config.get("official_sellers", [])
    official_domains = config.get("official_domains", [])
    pesos = config.get("pesos_relevancia", {
        "preco": 0.30, "autoridade_marca": 0.25, "presenca_ads": 0.20,
        "vendas_marketplace": 0.25, "kpis_redes_sociais": 0.0,
    })
    ativos_ads = {norm(k): v for k, v in load_json(args.ativos_ads_json, {}).items()} if args.ativos_ads_json else {}

    produtos = config.get("produtos_monitorados", []) + config.get("produtos_candidatos_manual", [])
    candidatos_manual = config.get("candidatos_concorrentes_manual", [])

    simulado = load_json(args.simulate_raw, {}) if args.simulate_raw else None
    token = None
    if simulado is None:
        token = get_token(config, args.token)
        if not token:
            print("ERRO: sem token Apify. Exporte APIFY_TOKEN, ou use --simulate-raw para testar.", file=sys.stderr)
            sys.exit(1)

    por_produto_pontuado = {}
    for prod in produtos:
        nome, termo = prod["nome"], prod["termo_busca_ml"]
        if simulado is not None:
            print(f"[SIMULAÇÃO] usando itens simulados para '{nome}'", file=sys.stderr)
            raw = simulado.get(nome, [])
        else:
            print(f"  buscando variantes/concorrentes de '{nome}' ({termo})...", file=sys.stderr)
            raw = buscar_raw(termo, token, args.per_produto)
        descoberta = descobrir_por_produto(raw, official_sellers)
        candidatos = compor_candidatos_produto(nome, descoberta, candidatos_manual)
        pontuados = [(c, pontuar(c, prod.get("preco_proprio"), ativos_ads, pesos)) for c in candidatos]
        por_produto_pontuado[nome] = pontuados

    dominios_auction = carregar_dominios_auction(args.auction_json, official_domains) if args.auction_json else {}
    candidatos_globais = compor_candidatos_globais(dominios_auction, candidatos_manual)
    globais_pontuado = [(c, pontuar(c, None, ativos_ads, pesos)) for c in candidatos_globais]

    write_xlsx(por_produto_pontuado, globais_pontuado, args.out)

    print(f"\nOK -> {args.out}", file=sys.stderr)
    for produto, pontuados in por_produto_pontuado.items():
        top = sorted(pontuados, key=lambda t: -(t[1]["score_final"] or -1))[:3]
        print(f"\nTop candidatos — {produto}:", file=sys.stderr)
        for cand, pont in top:
            score_txt = f"{pont['score_final']:.2f}" if pont["score_final"] is not None else "N/D"
            print(f"  • {cand.get('nome') or cand.get('seller')} — score {score_txt} "
                  f"(origem: {', '.join(cand.get('origem', []))})", file=sys.stderr)


if __name__ == "__main__":
    main()
